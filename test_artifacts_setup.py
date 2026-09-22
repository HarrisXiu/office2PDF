from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
p=Path('test_artifacts');p.mkdir(exist_ok=True)
d=Document();d.add_paragraph('Office2PDF conversion test');d.add_page_break();d.add_paragraph('Page two');d.save(p/'sample.docx')
w=Workbook();w.active.title='First';w.active.append(['Office2PDF',123]);w.create_sheet('Second').append(['Second sheet',456]);w.save(p/'sample.xlsx')
r=Presentation();r.slides.add_slide(r.slide_layouts[6]);r.slides.add_slide(r.slide_layouts[6]);r.save(p/'sample.pptx')
